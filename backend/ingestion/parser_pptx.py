import os
import subprocess
import logging
from pptx import Presentation
from ingestion.schemas import RawUnit
from app.config import settings

logger = logging.getLogger(__name__)

def parse_pptx(source_id: int, file_path: str) -> list[RawUnit]:
    units = []
    
    media_dir = os.path.join(settings.media_dir, str(source_id), "slides")
    os.makedirs(media_dir, exist_ok=True)
    
    # Render PDF via LibreOffice to get slide images
    pdf_path = os.path.join(settings.media_dir, str(source_id), f"{source_id}.pdf")
    try:
        # LibreOffice headless conversion
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", os.path.dirname(pdf_path), file_path],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL
        )
        
        # Then render PDF to PNGs using PyMuPDF
        import fitz
        doc = fitz.open(os.path.join(os.path.dirname(pdf_path), os.path.basename(file_path).replace(".pptx", ".pdf")))
        for page_num in range(len(doc)):
            page = doc[page_num]
            pix = page.get_pixmap()
            image_path = os.path.join(media_dir, f"slide_{page_num+1}.png")
            pix.save(image_path)
            
            # Extract slide as a figure unit
            units.append(RawUnit(
                text="",
                unit_type="figure",
                slide_no=page_num + 1,
                image_path=image_path
            ))
    except Exception as e:
        logger.error(f"Failed to render slides: {e}")
        
    # Extract text using python-pptx
    try:
        prs = Presentation(file_path)
        for i, slide in enumerate(prs.slides):
            slide_no = i + 1
            text_chunks = []
            
            # Extract text from shapes
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text_chunks.append(shape.text.strip())
                    
            # Extract speaker notes
            if slide.has_notes_slide:
                notes_slide = slide.notes_slide
                notes_text = notes_slide.notes_text_frame.text
                if notes_text.strip():
                    text_chunks.append("Speaker Notes: " + notes_text.strip())
                    
            text = "\n".join([t for t in text_chunks if t])
            if text:
                units.append(RawUnit(
                    text=text,
                    unit_type="text",
                    slide_no=slide_no
                ))
    except Exception as e:
        logger.error(f"Failed to parse PPTX text: {e}")

    # Return sorted by slide_no
    units.sort(key=lambda x: (x.slide_no, x.unit_type == "text"))
    return units
